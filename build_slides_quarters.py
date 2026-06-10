#!/usr/bin/env python3
"""RealWaterSimulator — architecture deck, organized by the four study quarters.

Follows the content/flow of QUARTER_1..4 md files:
  Q1 Application Host · Q2 Ocean Simulation · Q3 Rendering Engine · Q4 Water Interactions.
Same deep-ocean design language as the concept deck. Minimal text, native vector
diagrams, 16:9, built for a 5-minute talk.
"""
import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

# ---------------------------------------------------------------- palette
DEEP   = RGBColor(0x06, 0x12, 0x22)
DEEP2  = RGBColor(0x0B, 0x22, 0x3C)
MID    = RGBColor(0x12, 0x33, 0x52)
INK    = RGBColor(0xEA, 0xF2, 0xFA)
MUTE   = RGBColor(0x8B, 0xA3, 0xBC)
AQUA   = RGBColor(0x57, 0xC7, 0xE8)   # Q2 ocean
TEAL   = RGBColor(0x2A, 0xC0, 0xAE)   # Q4 interactions
GOLD   = RGBColor(0xF3, 0xC5, 0x6A)   # Q1 host
VIOLET = RGBColor(0x9D, 0x8C, 0xFF)   # Q3 rendering
BOXBG  = RGBColor(0x10, 0x2A, 0x47)
BOXLN  = RGBColor(0x24, 0x4B, 0x72)

EMU_IN = 914400
SW, SH = 13.333, 7.5

prs = Presentation()
prs.slide_width  = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]

# ---------------------------------------------------------------- helpers
def slide():
    return prs.slides.add_slide(BLANK)

def _grad(fill, c1, c2, angle):
    fill.gradient()
    try: fill.gradient_angle = angle
    except Exception: pass
    st = fill.gradient_stops
    st[0].position = 0.0; st[0].color.rgb = c1
    st[1].position = 1.0; st[1].color.rgb = c2

def bg(s, c1=DEEP2, c2=DEEP, angle=90):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.line.fill.background(); r.shadow.inherit = False
    _grad(r.fill, c1, c2, angle)
    return r

def text(s, txt, x, y, w, h, size, color=INK, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri Light",
         spacing=None, line_spacing=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, ln in enumerate(txt.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing: p.line_spacing = line_spacing
        r = p.add_run(); r.text = ln
        f = r.font; f.size = Pt(size); f.bold = bold; f.italic = italic
        f.name = font; f.color.rgb = color
        if spacing is not None:
            r.font._rPr.set('spc', str(int(spacing * 100)))
    return tb

def kicker(s, txt, color=AQUA, x=0.92, y=0.62):
    text(s, txt.upper(), x, y, 11, 0.4, 13, color, bold=True, font="Calibri", spacing=2.5)

def title(s, txt, x=0.92, y=1.02, w=11.6, size=33, color=INK):
    text(s, txt, x, y, w, 1.0, size, color, font="Calibri Light")

def rule(s, x, y, w, color=AQUA, weight=2.2):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x+w), Inches(y))
    ln.line.color.rgb = color; ln.line.width = Pt(weight); ln.shadow.inherit = False
    return ln

def box(s, x, y, w, h, fill=BOXBG, line=BOXLN, lw=1.25, round=True, radius=0.10):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    if round:
        try: shp.adjustments[0] = radius
        except Exception: pass
    return shp

def topbar(s, x, y, w, color, h=0.12):
    return box(s, x, y, w, h, fill=color, line=None, round=False)

def chevron(s, x, y, color=AQUA, w=0.20, h=0.34):
    a = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background(); a.shadow.inherit = False
    return a

def down_arrow(s, x, y, color=AQUA, w=0.34, h=0.32):
    a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background(); a.shadow.inherit = False
    return a

def connect(s, x1, y1, x2, y2, color=MUTE, weight=1.4):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color; ln.line.width = Pt(weight); ln.shadow.inherit = False
    return ln

def flowbox(s, x, y, w, h, head, sub, num=None, head_color=AQUA):
    box(s, x, y, w, h); pad = 0.18
    if num is not None:
        text(s, num, x+pad, y+0.12, w-2*pad, 0.4, 12, MUTE, bold=True, font="Calibri")
    text(s, head, x+pad, y+0.40, w-2*pad, 0.5, 15.5, head_color, bold=True, font="Calibri")
    text(s, sub, x+pad, y+0.90, w-2*pad, h-0.95, 11.5, MUTE, font="Calibri Light", line_spacing=1.0)

def dot(s, cx, cy, d, color):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx-d/2), Inches(cy-d/2), Inches(d), Inches(d))
    o.fill.solid(); o.fill.fore_color.rgb = color
    o.line.fill.background(); o.shadow.inherit = False
    return o

def sine(s, x, y, w, amp, cycles, color=AQUA, weight=2.5, n=140, phase=0.0):
    pts = []
    for i in range(n+1):
        t = i/n
        px = (x + t*w)*EMU_IN
        py = (y - amp*math.sin(phase + t*cycles*2*math.pi))*EMU_IN
        pts.append((Emu(int(px)), Emu(int(py))))
    fb = s.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
    fb.add_line_segments(pts[1:], close=False)
    shp = fb.convert_to_shape()
    shp.fill.background(); shp.line.color.rgb = color; shp.line.width = Pt(weight)
    shp.shadow.inherit = False
    return shp

def page(s, n):
    text(s, f"{n:02d}", SW-1.05, SH-0.62, 0.7, 0.4, 11, MUTE, align=PP_ALIGN.RIGHT, font="Calibri")
    text(s, "RealWaterSimulator  ·  architecture", 0.92, SH-0.62, 7, 0.4, 11, MUTE, font="Calibri")

Q = [("Q1", "Application Host", GOLD), ("Q2", "Ocean Simulation", AQUA),
     ("Q3", "Rendering Engine", VIOLET), ("Q4", "Water Interactions", TEAL)]

# ================================================================ SLIDE 1 — Title
s = slide(); bg(s, DEEP2, DEEP, 90)
sine(s, 0.0, 5.55, SW, 0.16, 3.0, color=RGBColor(0x1E,0x55,0x7A), weight=2.0, phase=0.4)
sine(s, 0.0, 5.95, SW, 0.22, 2.3, color=RGBColor(0x2A,0x6E,0x9C), weight=2.2, phase=1.1)
sine(s, 0.0, 6.42, SW, 0.30, 1.7, color=AQUA, weight=2.8, phase=0.0)
text(s, "Real-Time Water Simulator", 0.92, 2.15, 11.5, 1.3, 52, INK)
rule(s, 1.0, 3.32, 2.1, AQUA, 2.5)
text(s, "Inside the build — a four-part architecture", 0.92, 3.5, 11, 0.7, 22, AQUA)
# four quarter chips with colored dots
qx = 0.97
for i,(q,name,c) in enumerate(Q):
    yy = 4.45
    dot(s, qx+0.08, yy+0.14, 0.16, c)
    text(s, f"{q}", qx+0.32, yy-0.05, 0.6, 0.4, 14, c, bold=True, font="Calibri")
    text(s, name, qx+0.74, yy-0.05, 2.5, 0.4, 14, INK, font="Calibri Light")
    qx += 3.02
text(s, "C++  ·  OpenGL 4.6 compute  ·  CMake, fully vendored  ·  100+ FPS", 0.95, 6.9, 11, 0.4, 13, MUTE, font="Calibri")

# ================================================================ SLIDE 2 — Architecture overview
s = slide(); bg(s); kicker(s, "Architecture", INK)
title(s, "Four quarters, one frame loop")
rule(s, 0.95, 1.78, 1.4, INK)
# hub: Q1
hub_w, hub_h, hub_x, hub_y = 6.7, 1.25, (SW-6.7)/2, 2.1
box(s, hub_x, hub_y, hub_w, hub_h, fill=MID, line=GOLD, lw=1.75)
topbar(s, hub_x, hub_y, hub_w, GOLD)
text(s, "Q1 · Application Host", hub_x+0.3, hub_y+0.32, hub_w-0.6, 0.5, 20, GOLD, bold=True, font="Calibri")
text(s, "the frame loop that creates, orders, and wires the other three", hub_x+0.3, hub_y+0.78, hub_w-0.6, 0.4, 14, MUTE, font="Calibri Light")
# spokes
spokes = [("Q2 · Ocean Simulation", "the moving, lit water surface", AQUA),
          ("Q3 · Rendering Engine", "window, camera, shaders, models", VIOLET),
          ("Q4 · Water Interactions", "buoyancy, rain, wakes, rock", TEAL)]
sw_, sh_, sy = 3.6, 1.7, 4.55
gap = (11.45 - 3*sw_) / 2
sx0 = 0.95
for i,(h,sub,c) in enumerate(spokes):
    x = sx0 + i*(sw_+gap)
    box(s, x, sy, sw_, sh_, fill=BOXBG, line=BOXLN)
    topbar(s, x, sy, sw_, c)
    text(s, h, x+0.24, sy+0.34, sw_-0.48, 0.5, 16.5, c, bold=True, font="Calibri")
    text(s, sub, x+0.24, sy+0.92, sw_-0.48, 0.6, 13.5, MUTE, font="Calibri Light", line_spacing=1.05)
    connect(s, x+sw_/2, hub_y+hub_h, x+sw_/2, sy, color=BOXLN, weight=1.6)
text(s, "Q1 is the hub; the other three are spokes it instantiates and calls each frame.",
     0.95, 6.55, 11.5, 0.5, 15, INK, italic=True)
page(s, 2)

# ================================================================ SLIDE 3 — Q1
s = slide(); bg(s); kicker(s, "Q1 · Application Host", GOLD)
title(s, "One loop, advanced in fixed steps")
rule(s, 0.95, 1.78, 1.4, GOLD)
stages = [
    ("INPUT","keyboard + mouse, UI panel"),
    ("PHYSICS","fixed 1/60 s:  ocean · wakes · steering"),
    ("READBACK","wave heights, one frame stale"),
    ("BUOYANCY","float + tilt each boat"),
    ("RENDER","ocean · objects · sky · rain"),
]
bw, bh, gap, x0, y0 = 2.18, 1.95, 0.20, 0.95, 2.3
for i,(h,sub) in enumerate(stages):
    x = x0 + i*(bw+gap)
    flowbox(s, x, y0, bw, bh, h, sub, num=f"{i+1}", head_color=GOLD)
    if i < len(stages)-1:
        chevron(s, x+bw, y0+bh/2-0.17, GOLD)
# two notes
ny = 4.85
box(s, 0.95, ny, 5.75, 1.5, fill=MID, line=BOXLN)
text(s, "Fixed-timestep accumulator", 1.2, ny+0.2, 5.3, 0.4, 16, GOLD, bold=True, font="Calibri")
text(s, "The sim advances in identical 1/60 s chunks — it looks the same at 30, 60, or 144 FPS instead of wobbling when the rate dips.",
     1.2, ny+0.62, 5.3, 0.8, 12.5, MUTE, font="Calibri Light", line_spacing=1.05)
box(s, 6.95, ny, 5.4, 1.5, fill=MID, line=BOXLN)
text(s, "Build system", 7.2, ny+0.2, 5.0, 0.4, 16, GOLD, bold=True, font="Calibri")
text(s, "CMake with every dependency vendored in the tree — GLFW, GLAD, GLM, cgltf, ImGui. One command builds on Linux and Windows.",
     7.2, ny+0.62, 5.0, 0.8, 12.5, MUTE, font="Calibri Light", line_spacing=1.05)
text(s, "Q1 simulates nothing itself — it choreographs the departments and hands each one's output to the next.",
     0.95, 6.6, 11.6, 0.4, 13, MUTE, italic=True, font="Calibri")
page(s, 3)

# ================================================================ SLIDE 4 — Q2
s = slide(); bg(s); kicker(s, "Q2 · Ocean Simulation", AQUA)
title(s, "Statistics become a surface")
rule(s, 0.95, 1.78, 1.4, AQUA)
# top: 3 flow boxes
fb = [("SPECTRUM","which waves exist, how tall — a real wind-wave model, frozen once"),
      ("PHASE","advance every wave in time by the dispersion relation"),
      ("IFFT","sum them all into a real surface — Stockham, rows then columns")]
bw, bh, gap, x0, y0 = 3.55, 1.55, 0.30, 0.95, 2.2
for i,(h,sub) in enumerate(fb):
    x = x0 + i*(bw+gap)
    flowbox(s, x, y0, bw, bh, h, sub, num=f"{i+1}", head_color=AQUA)
    if i < 2: chevron(s, x+bw+0.03, y0+bh/2-0.17, AQUA)
# down arrow to outputs
down_arrow(s, x0 + 2*(bw+gap) + bw/2 - 0.17, y0+bh+0.06, AQUA)
text(s, "three outputs the rest of the program leans on", 0.95, y0+bh+0.12, 7, 0.34, 12.5, MUTE, font="Calibri")
oy = y0+bh+0.55
outs = [("Displacement texture","where each point of water moves  (dx, h, dz)"),
        ("Normal + foam texture","which way it faces, and where it whitecaps"),
        ("CPU height mirror","so physics can ask: how high is the water here?")]
for i,(h,sub) in enumerate(outs):
    x = x0 + i*(bw+gap)
    box(s, x, oy, bw, 1.15, fill=MID, line=BOXLN)
    text(s, h, x+0.2, oy+0.16, bw-0.4, 0.4, 14.5, AQUA, bold=True, font="Calibri")
    text(s, sub, x+0.2, oy+0.58, bw-0.4, 0.5, 12, MUTE, font="Calibri Light", line_spacing=1.0)
text(s, "A player piano: the spectrum is the score; the GPU unfolds it into water every frame.   "
        "Stockham IFFT lifts the shared-memory cap → a true 512×512;  foam appears where the surface folds (Jacobian < 0).",
     0.95, 6.35, 11.7, 0.7, 13, INK, italic=True, font="Calibri Light", line_spacing=1.1)
page(s, 4)

# ================================================================ SLIDE 5 — Q3
s = slide(); bg(s); kicker(s, "Q3 · Rendering Engine", VIOLET)
title(s, "The studio rig everything plugs into")
rule(s, 0.95, 1.78, 1.4, VIOLET)
tools = [("Window","GL 4.6 core context — compute legal"),
         ("Camera","two angles → a view matrix"),
         ("Shader","the one CPU↔GPU adapter, used by all"),
         ("Mesh · Model","glTF props, node transforms baked in"),
         ("Texture","image files → GPU samplers")]
bw, gap, x0, y0, bh = 2.18, 0.20, 0.95, 2.45, 1.85
for i,(h,sub) in enumerate(tools):
    x = x0 + i*(bw+gap)
    box(s, x, y0, bw, bh, fill=BOXBG, line=BOXLN)
    topbar(s, x, y0, bw, VIOLET)
    text(s, h, x+0.18, y0+0.34, bw-0.36, 0.6, 15, INK, bold=False, font="Calibri")
    text(s, sub, x+0.18, y0+0.95, bw-0.36, 0.8, 11.5, MUTE, font="Calibri Light", line_spacing=1.05)
# shader notes
ny = 4.75
box(s, 0.95, ny, 5.75, 1.5, fill=MID, line=BOXLN)
text(s, "object.frag — solid props", 1.2, ny+0.2, 5.3, 0.4, 15.5, VIOLET, bold=True, font="Calibri")
text(s, "PBR lit by sun + sky reflection. The TBN frame is rebuilt from screen-space derivatives, so models need no tangent data.",
     1.2, ny+0.62, 5.3, 0.8, 12.5, MUTE, font="Calibri Light", line_spacing=1.05)
box(s, 6.95, ny, 5.4, 1.5, fill=MID, line=BOXLN)
text(s, "skybox — the backdrop", 7.2, ny+0.2, 5.0, 0.4, 15.5, VIOLET, bold=True, font="Calibri")
text(s, "Drawn last, pinned to the far plane with the .xyww depth trick, camera-centred. The same cube is reused as the reflection.",
     7.2, ny+0.62, 5.0, 0.8, 12.5, MUTE, font="Calibri Light", line_spacing=1.05)
text(s, "Shared by every quarter — Q2 and Q4 compile their compute passes through this same Shader class.",
     0.95, 6.55, 11.6, 0.4, 13, MUTE, italic=True, font="Calibri")
page(s, 5)

# ================================================================ SLIDE 6 — Q4
s = slide(); bg(s); kicker(s, "Q4 · Water Interactions", TEAL)
title(s, "What makes the inert ocean feel alive")
rule(s, 0.95, 1.78, 1.4, TEAL)
cells = [
    ("Buoyancy", "Boats float on springs, not forces — 15 hull facets sample the wave; average height → heave, surface slope → pitch and roll."),
    ("GPU rain", "120,000 drops in one buffer. One compute dispatch moves them all; two attributeless draws render streaks and splashes."),
    ("Wake field", "A literal pond simulation — the Verlet wave equation on a 256² grid. Moving boats inject a stern ripple that spreads and fades."),
    ("Procedural rock", "A geometry recipe — icosphere, subdivided, then pushed by layered noise into a boulder. A fallback when the glTF rock is missing."),
]
cw, ch, gx, gy = 5.6, 1.7, 0.95, 2.4
gap_x, gap_y = 0.25, 0.3
for i,(h,sub) in enumerate(cells):
    r, c = divmod(i, 2)
    x = gx + c*(cw+gap_x); y = gy + r*(ch+gap_y)
    box(s, x, y, cw, ch, fill=BOXBG, line=BOXLN)
    dot(s, x+0.42, y+0.45, 0.18, TEAL)
    text(s, h, x+0.72, y+0.22, cw-1.0, 0.5, 17, INK, font="Calibri")
    text(s, sub, x+0.3, y+0.78, cw-0.6, 0.85, 12.5, MUTE, font="Calibri Light", line_spacing=1.05)
text(s, "Q4 owns no water render — it feeds height and ripples into Q2's shader and is ticked by Q1's loop.",
     0.95, 6.5, 11.6, 0.4, 13.5, MUTE, italic=True, font="Calibri")
page(s, 6)

# ================================================================ SLIDE 7 — The contract
s = slide(); bg(s); kicker(s, "How the quarters connect", INK)
title(s, "One shared GPU contract")
rule(s, 0.95, 1.78, 1.4, INK)
text(s, "Every quarter agrees on the same texture-unit budget — collisions would corrupt each other's bindings.",
     0.95, 2.05, 11.6, 0.4, 14.5, MUTE, font="Calibri Light")
# 8 unit chips
units = [("0","sky cube", GOLD), ("1","displace", AQUA), ("2","normal", AQUA),
         ("3","wake", TEAL), ("4","albedo", VIOLET), ("5","normal", VIOLET),
         ("6","rough", VIOLET), ("7","AO", VIOLET)]
cw, gap, x0, y0, ch = 1.30, 0.14, 0.95, 2.95, 1.25
for i,(num,lab,c) in enumerate(units):
    x = x0 + i*(cw+gap)
    box(s, x, y0, cw, ch, fill=BOXBG, line=c, lw=1.5)
    text(s, num, x, y0+0.16, cw, 0.55, 26, c, bold=True, align=PP_ALIGN.CENTER, font="Calibri Light")
    text(s, lab, x, y0+0.78, cw, 0.35, 11.5, MUTE, align=PP_ALIGN.CENTER, font="Calibri")
# group brackets / labels
text(s, "ocean  (units 0–3)", x0, y0+ch+0.12, 4*cw+3*gap, 0.35, 12.5, AQUA, align=PP_ALIGN.CENTER, font="Calibri", bold=True)
gx2 = x0 + 4*(cw+gap)
text(s, "object materials  (units 4–7)", gx2, y0+ch+0.12, 4*cw+3*gap, 0.35, 12.5, VIOLET, align=PP_ALIGN.CENTER, font="Calibri", bold=True)
# SSBO chip
ssy = y0+ch+0.7
box(s, 0.95, ssy, 5.6, 0.9, fill=MID, line=TEAL, lw=1.5)
text(s, "SSBO binding 4", 1.2, ssy+0.16, 3, 0.4, 15, TEAL, bold=True, font="Calibri")
text(s, "rain ripple rings, read by the water shader", 1.2, ssy+0.55, 5.2, 0.3, 12, MUTE, font="Calibri")
box(s, 6.75, ssy, 5.65, 0.9, fill=MID, line=BOXLN)
text(s, "Q2 produces textures · Q1 wires them · Q3 supplies the GL toolkit + skybox · Q4 feeds wake + ripples back and reads heights.",
     7.0, ssy+0.12, 5.2, 0.7, 12, INK, font="Calibri Light", line_spacing=1.05, anchor=MSO_ANCHOR.MIDDLE)
page(s, 7)

# ================================================================ SLIDE 8 — Demo / close
s = slide(); bg(s, DEEP2, DEEP, 90)
sine(s, 0.0, 5.7, SW, 0.18, 2.6, color=RGBColor(0x1E,0x55,0x7A), weight=2.0, phase=0.4)
sine(s, 0.0, 6.15, SW, 0.26, 2.0, color=RGBColor(0x2A,0x6E,0x9C), weight=2.3, phase=1.1)
sine(s, 0.0, 6.6, SW, 0.34, 1.5, color=AQUA, weight=3.0, phase=0.0)
text(s, "Demo", 0.92, 2.35, 8, 1.4, 60, INK)
rule(s, 1.0, 3.7, 2.1, AQUA, 2.5)
text(s, "All four quarters, in one live frame.", 0.95, 3.92, 11.5, 0.7, 22, AQUA)
# four dots inline
qx = 0.97
for i,(q,name,c) in enumerate(Q):
    yy = 4.85
    dot(s, qx+0.08, yy+0.13, 0.15, c)
    text(s, name, qx+0.3, yy-0.05, 2.7, 0.4, 13.5, INK, font="Calibri Light")
    qx += 3.02
text(s, "Drive the boat · turn calm into storm · sweep the sun · make it rain.", 0.95, 5.95, 11, 0.5, 14, MUTE, italic=True, font="Calibri")

prs.save("/home/bora/Projects/learning/RealWaterSimulatorOpenGL/RealWaterSimulator_Architecture.pptx")
print("saved RealWaterSimulator_Architecture.pptx with", len(prs.slides._sldIdLst), "slides")
